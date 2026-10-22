import base64
import mimetypes
from io import BytesIO

import frappe
from frappe.utils import (
    add_days, cint, escape_html, fmt_money, formatdate, get_first_day,
    getdate, money_in_words, now_datetime, today,
)
from frappe.utils.pdf import get_pdf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def create_monthly_petty_cash_whish():
    get_or_create_monthly_petty_cash_whish(get_first_day(today()))


def get_or_create_monthly_petty_cash_whish(month):
    month = get_first_day(month)

    existing = frappe.db.exists(
        "Petty Cash Whish",
        {
            "month_and_year": month
        }
    )

    if existing:
        return existing

    whish = frappe.get_doc({
        "doctype": "Petty Cash Whish",
        "month_and_year": month
    })

    try:
        whish.insert(ignore_permissions=True)
    except frappe.DuplicateEntryError:
        # Another worker created the same uniquely named monthly record.
        return frappe.db.get_value(
            "Petty Cash Whish", {"month_and_year": month}, "name"
        )

    return whish.name


@frappe.whitelist()
def send_monthly_reminders(force=0):
    force = cint(force)
    if force and not ({"LSA Finance Approver", "System Manager"} & set(frappe.get_roles())):
        frappe.throw("Only Finance or System Manager may send reminders.", frappe.PermissionError)

    settings = frappe.get_single("Petty Cash Settings")
    current_date = getdate(today())
    month_key = current_date.strftime("%Y-%m")

    if not force:
        if not settings.enable_monthly_reminders or current_date.day != cint(settings.reminder_day):
            return 0

    sent = 0
    configurations = frappe.get_all(
        "Petty Cash Configuration",
        fields=["name", "center_officer", "reminder_email", "last_reminder_month"],
    )
    for config in configurations:
        if not force and config.last_reminder_month == month_key:
            continue
        if frappe.db.exists(
            "Petty Cash Settlement",
            {"center_officer": config.center_officer, "settlement_month": month_key, "docstatus": ["<", 2]},
        ):
            continue

        recipient = config.reminder_email or _employee_email(config.center_officer)
        if not recipient:
            frappe.log_error(
                f"No reminder email is configured for Employee {config.center_officer}",
                "Petty Cash Reminder Skipped",
            )
            continue

        frappe.sendmail(
            recipients=[recipient],
            subject=settings.reminder_subject,
            message=f"<p>{frappe.utils.escape_html(settings.reminder_message)}</p>",
        )
        frappe.db.set_value(
            "Petty Cash Configuration",
            config.name,
            {"last_reminder_month": month_key, "reminder_sent_on": now_datetime()},
            update_modified=False,
        )
        sent += 1

    return sent


def _employee_email(employee):
    user = frappe.db.get_value("Employee", employee, "user_id")
    if not user:
        return None
    return frappe.db.get_value("User", user, "email") or user


def generate_whish_excel_attachment(payment_date=None, whish_name=None):
    if not whish_name and payment_date:
        whish_month = get_first_day(payment_date)
        whish_name = frappe.db.get_value(
            "Petty Cash Whish", {"month_and_year": whish_month}, "name"
        )
    if not whish_name:
        frappe.throw("The monthly Petty Cash Whish document does not exist.")

    whish = frappe.get_doc("Petty Cash Whish", whish_name)
    rows = [[
        "Number",
        "NAME",
        "Phone Number",
        "Whish ID",
        "Salary received by the employee ( NET)",
        "Currency",
    ]]
    for number, row in enumerate(whish.employees, start=1):
        rows.append([
            number,
            _safe_excel_value(row.employee_name),
            _safe_excel_value(row.phone_number),
            _safe_excel_value(row.whish_id),
            row.salary_received_net,
            _safe_excel_value(row.currency),
        ])

    content = _make_whish_excel(rows)
    file_version = now_datetime().strftime("%Y%m%d-%H%M%S-%f")
    attachment_name = f"{whish.name}-v{file_version}.xlsx"
    old_files = frappe.get_all(
        "File",
        filters={
            "attached_to_doctype": "Petty Cash Whish",
            "attached_to_name": whish.name,
            "file_name": ["like", "%.xlsx"],
        },
        pluck="name",
    )
    for old_file in old_files:
        frappe.delete_doc("File", old_file, ignore_permissions=True)

    file_doc = frappe.get_doc({
        "doctype": "File",
        "file_name": attachment_name,
        "content": content,
        "is_private": 1,
        "attached_to_doctype": "Petty Cash Whish",
        "attached_to_name": whish.name,
    })
    file_doc.save(ignore_permissions=True)
    whish.db_set("excel_file", file_doc.file_url)
    return file_doc.name


def generate_whish_pdf_attachment(whish_name):
    whish = frappe.get_doc("Petty Cash Whish", whish_name)
    context = _get_whish_letter_context(whish)
    content = get_pdf(_build_whish_letter_html(context))
    version = now_datetime().strftime("%Y%m%d-%H%M%S-%f")
    file_name = f"{whish.name}-payment-letter-v{version}.pdf"

    old_files = frappe.get_all(
        "File",
        filters={
            "attached_to_doctype": "Petty Cash Whish",
            "attached_to_name": whish.name,
            "file_name": ["like", "%-payment-letter-%.pdf"],
        },
        pluck="name",
    )
    for old_file in old_files:
        frappe.delete_doc("File", old_file, ignore_permissions=True)

    file_doc = frappe.get_doc({
        "doctype": "File",
        "file_name": file_name,
        "content": content,
        "is_private": 1,
        "attached_to_doctype": "Petty Cash Whish",
        "attached_to_name": whish.name,
    })
    file_doc.save(ignore_permissions=True)
    whish.db_set("pdf_file", file_doc.file_url)
    return file_doc.name


def generate_whish_attachments(whish_name):
    return {
        "excel": generate_whish_excel_attachment(whish_name=whish_name),
        "pdf": generate_whish_pdf_attachment(whish_name),
    }


@frappe.whitelist()
def get_whish_pdf_preview(whish_name):
    if not frappe.has_permission("Petty Cash Whish", "read", whish_name):
        frappe.throw("You do not have permission to view this Whish batch.", frappe.PermissionError)
    file_name = generate_whish_pdf_attachment(whish_name)
    return frappe.db.get_value("File", file_name, "file_url")


# Backward-compatible target for jobs queued before this change. It deliberately
# refreshes the attachment only and never sends an email.
def generate_and_email_whish_excel(payment_date):
    return generate_whish_excel_attachment(payment_date=payment_date)


@frappe.whitelist()
def send_whish_excel_now(whish_name):
    if not ({"LSA Finance Approver", "System Manager"} & set(frappe.get_roles())):
        frappe.throw("Only Finance or System Manager may send the Whish Excel.", frappe.PermissionError)
    return _send_whish_excel(whish_name)


def send_scheduled_whish_excel():
    """Retained as a safe no-op for workers that cached the old scheduler hook."""
    return None


def _send_whish_excel(whish_name):
    whish = frappe.get_doc("Petty Cash Whish", whish_name)
    if not whish.employees:
        frappe.throw("Add at least one approved settlement before sending the Whish batch.")

    settings = frappe.get_single("Petty Cash Settings")
    if not settings.whish_email:
        frappe.throw("Whish Email must be configured in Petty Cash Settings.")

    attachments = generate_whish_attachments(whish_name)
    excel_file = frappe.get_doc("File", attachments["excel"])
    pdf_file = frappe.get_doc("File", attachments["pdf"])

    for row in whish.employees:
        settlement = frappe.get_doc("Petty Cash Settlement", row.settlement)
        if settlement.workflow_state != "Approved for Whish Payment":
            frappe.throw(f"Settlement {settlement.name} is not approved for Whish payment.")
        settlement.create_draft_whish_journal_entry()

    context = _get_whish_letter_context(whish)
    frappe.sendmail(
        recipients=[settings.whish_email],
        subject=_render_template(settings.whish_email_subject, context),
        message="<p>" + _render_template(settings.whish_email_message, context, html=True) + "</p>",
        attachments=[
            {"fname": excel_file.file_name, "fcontent": excel_file.get_content()},
            {"fname": pdf_file.file_name, "fcontent": pdf_file.get_content()},
        ],
    )
    whish.db_set("excel_email_sent_on", now_datetime())
    return attachments


def _get_whish_letter_context(whish):
    currencies = {row.currency for row in whish.employees if row.currency}
    if len(currencies) > 1:
        frappe.throw("All employees in a Whish batch must use the same currency.")
    currency = next(iter(currencies), "USD")
    total = sum(row.salary_received_net or 0 for row in whish.employees)
    month = getdate(whish.month_and_year)
    period_end = add_days(get_first_day(month), 19)
    period_start = add_days(get_first_day(month), -10)
    reference = f"LSA-PC-{month:%Y-%m}"
    settings = frappe.get_single("Petty Cash Settings")

    company = None
    if whish.employees:
        company = frappe.db.get_value(
            "Petty Cash Settlement", whish.employees[0].settlement, "company"
        )
    logo = settings.letter_logo
    if not logo and company:
        logo = frappe.db.get_value("Company", company, "company_logo")

    return {
        "amount": fmt_money(total, precision=2, currency=currency),
        "amount_in_words": money_in_words(total, currency),
        "beneficiaries": len(whish.employees),
        "currency": currency,
        "date": formatdate(today()),
        "month": month.strftime("%B %Y"),
        "period": f"{formatdate(period_start)} to {formatdate(period_end)}",
        "reference": reference,
        "settings": settings,
        "logo": _file_as_data_uri(logo),
        "signature_1": _file_as_data_uri(settings.signature_1_image),
        "signature_2": _file_as_data_uri(settings.signature_2_image),
    }


def _render_template(value, context, html=False):
    rendered = value or ""
    for key in ("amount", "amount_in_words", "beneficiaries", "currency", "date", "month", "period", "reference"):
        rendered = rendered.replace("{" + key + "}", str(context[key]))
    if html:
        return escape_html(rendered).replace("\n", "<br>")
    return rendered


def _file_as_data_uri(file_url):
    if not file_url:
        return ""
    file_name = frappe.db.get_value("File", {"file_url": file_url}, "name")
    if not file_name:
        return ""
    file_doc = frappe.get_doc("File", file_name)
    mime = mimetypes.guess_type(file_doc.file_name)[0] or "image/png"
    encoded = base64.b64encode(file_doc.get_content()).decode()
    return f"data:{mime};base64,{encoded}"


def _signature_html(number, context):
    settings = context["settings"]
    image = context[f"signature_{number}"]
    name = escape_html(settings.get(f"signature_{number}_name") or "")
    role = escape_html(settings.get(f"signature_{number}_role") or "")
    date_value = context["date"] if settings.date_under_signature == "Current date" else "________________"
    image_html = f'<img src="{image}" class="signature-image">' if image else '<div class="signature-line"></div>'
    return f"""
        <div class="signature-box">
            <div class="signature-heading">AUTHORIZED SIGNATORY {number}</div>
            {image_html}
            <div><b>Name:</b> {name}</div>
            <div><b>Role:</b> {role}</div>
            <div><b>Date:</b> {escape_html(date_value)}</div>
        </div>
    """


def _build_whish_letter_html(context):
    settings = context["settings"]
    logo = f'<img src="{context["logo"]}" class="logo">' if context["logo"] else ""
    letter_text = _render_template(settings.letter_text, context, html=True)
    return f"""<!doctype html><html><head><meta charset="utf-8"><style>
        @page {{ size: A4; margin: 18mm; }}
        body {{ font-family: Arial, sans-serif; font-size: 11px; color: #111; line-height: 1.5; }}
        .header {{ display:flex; justify-content:space-between; align-items:flex-start; }}
        .logo {{ max-width:70px; max-height:70px; }}
        .meta {{ text-align:right; font-size:10px; }}
        h2 {{ font-size:15px; text-decoration:underline; margin:14px 0 8px; }}
        .letter {{ margin:12px 0; }}
        table {{ width:100%; border-collapse:collapse; margin:12px 0 18px; }}
        td {{ border:1px solid #777; padding:6px; }} td:first-child {{ width:34%; background:#e8eef7; font-weight:bold; }}
        .signatures {{ display:flex; gap:24px; margin-top:18px; }}
        .signature-box {{ width:48%; }} .signature-heading {{ color:#666; margin-bottom:16px; }}
        .signature-line {{ height:42px; border-bottom:1px solid #555; margin-bottom:6px; }}
        .signature-image {{ max-height:52px; max-width:180px; display:block; margin-bottom:6px; }}
    </style></head><body>
        <div class="header"><div>{logo}</div><div class="meta"><b>Date:</b> {escape_html(context['date'])}<br><b>Reference:</b> {escape_html(context['reference'])}</div></div>
        <p><b>To:</b> {escape_html(settings.addressed_to)}</p>
        <h2>{escape_html(settings.letter_title)}</h2>
        <div class="letter">{letter_text}</div>
        <table>
            <tr><td>Total amount</td><td>{escape_html(context['amount'])}</td></tr>
            <tr><td>Amount in words</td><td>{escape_html(context['amount_in_words'])}</td></tr>
            <tr><td>Payment period</td><td>{escape_html(context['period'])}</td></tr>
            <tr><td>Payroll month</td><td>{escape_html(context['month'])}</td></tr>
            <tr><td>Beneficiaries</td><td>{context['beneficiaries']}</td></tr>
        </table>
        <div class="signatures">{_signature_html(1, context)}{_signature_html(2, context)}</div>
    </body></html>"""


def _safe_excel_value(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value or ""


def _make_whish_excel(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Whish"

    for row in rows:
        sheet.append(row)

    header_fill = PatternFill("solid", fgColor="9DC3E6")
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in sheet.iter_rows():
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal="center", vertical="center")

    for cell in sheet["E"][1:]:
        cell.number_format = "#,##0.00"

    widths = {"A": 12, "B": 30, "C": 18, "D": 18, "E": 38, "F": 14}
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    sheet.freeze_panes = "A2"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
